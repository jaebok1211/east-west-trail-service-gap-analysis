from __future__ import annotations

import subprocess
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=85, bottom=65, end=85) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_title(doc: Document, title: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(15)
    if subtitle:
        rr = p.add_run("\n" + subtitle)
        rr.font.size = Pt(8.7)


def add_body(doc: Document, text: str, size: float = 9.0, after: float = 3.0) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)


def add_caption(doc: Document, text: str, source: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(7.6)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p2.add_run(source)
    rr.font.size = Pt(6.8)


def add_source(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.size = Pt(6.9)


def add_table(doc: Document, headers, rows, widths=None, font=7.6, header_fill="DDEBF7"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = str(h)
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(font)
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = "" if value is None else str(value)
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(font)
    if widths:
        for row in t.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Cm(width)
    return t


def add_callout(doc: Document, heading: str, text: str, fill="EAF2F8") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=100, bottom=100, start=130, end=130)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(heading)
    r.bold = True
    r.font.size = Pt(8.5)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    rr = p2.add_run(text)
    rr.font.size = Pt(8.2)


def build(csv_dir: Path, fig_dir: Path, inputs_dir: Path, report_dir: Path) -> tuple[Path, Path]:
    report_dir.mkdir(parents=True, exist_ok=True)
    overall = pd.read_csv(csv_dir / "01_survey_overall.csv")
    demo = pd.read_csv(csv_dir / "02_survey_demographic_rates.csv")
    odds = pd.read_csv(csv_dir / "03_logistic_odds_ratios.csv")
    consumption = pd.read_csv(csv_dir / "06_consumption_summary.csv")
    lodging = pd.read_csv(csv_dir / "07_overnight_lodging_types.csv")
    region = pd.read_csv(csv_dir / "08_region_tourism_summary.csv")
    priority = pd.read_csv(csv_dir / "14_priority_table_final.csv")
    refs = pd.read_csv(inputs_dir / "reference_list.csv")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.25)
    sec.bottom_margin = Cm(1.18)
    sec.left_margin = Cm(1.55)
    sec.right_margin = Cm(1.55)
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK KR"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
    normal.font.size = Pt(9.0)

    header = sec.header.paragraphs[0]
    header.text = "동서트레일 공개구간의 체류서비스 공백 분석을 통한 거점 보완 우선순위 도출"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs:
        r.font.size = Pt(7.2)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("- ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
    footer.add_run(" -")

    # 1p
    add_title(doc, "1) 추진배경 및 필요성", "장거리 숲길의 물리적 연결을 지역 체류와 소비로 이어가기 위한 사전 진단")
    add_body(doc, "산림휴양은 숲을 방문하는 활동을 넘어 인근 지역의 숙박·식사·체험과 연결될 때 농·산촌의 소득 기반으로 확장될 수 있다. 동서트레일은 충남 태안에서 경북 울진까지 55개 구간, 849km를 연결하는 장거리 숲길로, 산림청은 2027년 전 구간 개통을 목표로 공개구간과 운영 프로그램을 단계적으로 확대하고 있다.[1-4]")
    add_body(doc, "사업 방향도 노선 개통에 머물지 않는다. 거점안내소와 숲길쉼터, 거점마을 협력, 예약·완주 인증, 지역 소득사업을 함께 추진하고 있어, 향후 정책의 쟁점은 ‘어느 구간에 어떤 기능을 먼저 보완할 것인가’로 이동한다.[3-4]")
    add_body(doc, "기존 연구는 도보여행길이 마을 안의 숙박·식사·체험과 연결되지 않으면 지역경제 효과가 제한될 수 있다고 보았다. 지리산길 이용객 연구에서도 숙박비와 식비가 주요 소비항목으로 나타났으며, 산촌마을은 획일적 시설 공급보다 지역 여건에 맞춘 운영 프로그램이 필요하다고 제시됐다.[11-13]")
    add_body(doc, "동서트레일의 구간은 거리, 소요시간, 난이도, 시·종점 입지와 주변 서비스가 다르다. 어떤 구간은 숙박이 아니라 식사 기능이 부족하고, 다른 구간은 상업숙박은 적지만 공식 대피소와 민박을 연계할 수 있다. 전체 시설 수만 비교하면 이러한 차이를 놓치며, 모든 거점에 동일한 시설을 배치하면 투자 효율도 낮아질 수 있다.")
    add_callout(doc, "연구 질문", "이동부담이 큰 공개구간 가운데 숙박·식음·보급·교통 중 어떤 기능이 실제로 비어 있으며, 구간별로 어떤 운영방식을 우선 적용해야 하는가?")
    add_body(doc, "이에 2025년 시범운영 이력이 있는 17개 구간을 대상으로 숙박형 산림활동의 수요 특성, 공식 노선의 이동부담, 양 끝점의 서비스 접근성, 시·군의 체류시장 여건을 각각 분석했다. 결과는 실제 이용자의 소비효과를 추정하는 자료가 아니라, 공개구간 확대 전에 체류서비스 공백과 보완 우선순위를 진단하는 정책자료로 해석한다.")
    add_source(doc, "주: 대괄호 번호는 10쪽 참고자료 번호와 연결된다.")
    doc.add_page_break()

    # 2p
    add_title(doc, "2) 아이디어 기획 세부 내용", "구간의 이동부담과 양 끝점 서비스 접근성을 결합한 거점기능 배분모델")
    add_body(doc, "본 아이디어는 시설이 많은 구간과 적은 구간을 단순 비교하지 않는다. 하루 이동을 마친 뒤 이용자가 실제로 도착하는 종점과 다음 이동을 시작하는 지점에서 숙박, 식사, 보급, 귀환 기능을 이용할 수 있는지를 점검하고, 이동부담에 비해 기능이 부족한 구간을 우선보완 대상으로 선정한다.")
    add_table(doc, ["1단계: 수요 특성", "2단계: 구간 공백", "3단계: 운영유형"], [[
        "숙박형 활동의 연령·활동·목적·소비 특성 확인",
        "공식 거리·시간·난이도와 시·종점 1km·3km 접근성 결합",
        "보급지원형·숙박전환형·기존자원연계형 등 기능별 처방"
    ]], widths=[5.5,5.5,5.5], font=8.0, header_fill="EAF2F8")
    add_body(doc, "분석단위가 다른 자료는 한 행으로 직접 병합하지 않았다. 개인 단위 조사는 숙박형 산림활동의 특성을, 구간 단위 자료는 서비스 공백을, 시·군 단위 관광자료는 사업 방식의 현실성을 설명한다. 세 결과는 최종 운영모델 배정 단계에서만 연결한다.")
    add_body(doc, "최종 산출물은 ① 17개 구간 우선순위 지도, ② 숙박·식음·보급·교통 기능별 공백표, ③ 상위구간의 실제 시설 위치와 1km·3km 접근권, ④ 구간별 운영유형으로 구성한다. 공식 대피소와 백패킹 지원기반은 상업숙박을 대체하는 동일 시설로 보지 않고, 숙박 공백을 해석하는 보조자료로 별도 검토한다.")
    add_table(doc, ["운영유형", "적용 조건", "우선 조치"], [
        ["보급지원형", "식사 또는 기초 보급 기능이 부족", "예약식사·도시락·식수·충전·영업정보 연계"],
        ["상업숙박·마을서비스 전환형", "상업숙박은 부족하나 대피소·민박 기반 존재", "민박·대피소·식사·샤워·짐보관 통합"],
        ["기본관리·기존자원연계형", "이동부담은 높지만 주요 서비스 기반 양호", "신규시설보다 예약·정보·지역업체 연결"],
        ["교통 추가확인 후보", "정류장 접근은 낮으나 운행정보 미확인", "노선·배차·막차 확인 후 셔틀 실증 여부 판단"],
    ], widths=[4.6,5.8,6.3], font=7.6)
    add_body(doc, "모델은 향후 개통구간에도 반복 적용할 수 있다. 공식 구간정보와 시설자료를 갱신하면 같은 코드로 접근성과 공백이 다시 산출되며, 실제 운영 이후 예약률·지역업체 이용·민원 자료를 반영해 우선순위를 조정할 수 있다.")
    doc.add_page_break()

    # 3p
    add_title(doc, "3) 데이터 분석 방법", "3-1. 자료 구성과 전처리")
    add_table(doc, ["분석 단위", "자료", "주요 변수", "산출물"], [
        ["개인", "2025 산림휴양복지활동조사[6]", "당일·숙박 경험, 활동, 목적, 동반, 소비, 인구특성", "숙박형 관련 특성·오즈비"],
        ["구간", "공식 GPX·코스정보[5]", "거리, 시간, 난이도, 시·종점", "체류·보급 필요도"],
        ["시설", "숙박·상가·버스 공공데이터[8-10]", "1km·3km 시설 수, 최근접 거리", "기능별 접근성"],
        ["시·군", "한국관광 데이터랩[7]", "방문, 숙박비율, 체류시간, 소비", "체류시장 여건 보조지표"],
    ], widths=[2.1,4.6,6.0,4.0], font=7.5)
    add_body(doc, "분석대상은 2025년 시범운영 구간 1~4, 9~12, 47~55의 17개다. 공식 GPX는 노선과 시·종점 좌표에 사용하고, 공식 거리·소요시간·난이도는 숲나들e 구간정보를 적용했다. 공식 진행방향과 반대인 9·55구간은 시·종점 기준으로 방향을 보정했다.[5]")
    add_body(doc, "숙박시설은 일반숙박과 농어촌민박을 영업 중인 사업장만 남기고, 사업장명과 주소가 같은 중복행을 제거했다. 좌표는 EPSG:5174에서 WGS84로 변환했다. 음식점과 보급시설은 충남·경북 상가정보에서 대상 시·군을 추출한 뒤, 음식점과 편의점·슈퍼·식료품점 등 실제 이동 중 이용 가능한 업종으로 구분했다.[9-10]")
    add_body(doc, "버스자료는 정류장 위치만 포함하므로 노선·배차를 직접 설명하지 않는다. 상위구간 중 정류장 접근거리가 긴 지점만 추가확인 후보로 표시했다.[8] 시설은 구간의 행정구역으로 선필터링하지 않고 전국 좌표에서 거리를 계산해, 시·군 경계 밖의 가까운 시설이 누락되지 않도록 했다.")
    add_body(doc, "원자료 공간분석 후 상위구간의 시설명·주소·영업상태와 업종을 다시 확인했다. 음식점과 기초 보급은 같은 범주가 아니므로 별도 집계했고, 검토 결과를 고정 입력표로 남겨 코드 실행 시 같은 최종 접근성표가 재생성되도록 했다.")
    add_callout(doc, "자료 단위 해석", "개인·구간·시·군 자료를 직접 연결해 ‘이 사람이 이 구간에서 얼마를 쓴다’고 추정하지 않는다. 각 자료가 설명하는 범위를 분리하고 정책 단계에서만 함께 해석한다.")
    add_source(doc, "분석기간: 설문 2025년 조사자료, 관광자료 2025.07~2026.06, 시설자료는 제출 패키지의 input_manifest.csv 참조.")
    doc.add_page_break()

    # 4p
    add_title(doc, "3) 데이터 분석 방법", "3-2. 지표 산정과 검증")
    add_body(doc, "개인 분석은 당일형만 경험한 응답자를 0, 숙박형 활동을 포함한 응답자를 1로 두고 표본가중치를 적용한 이항 일반화선형모형을 추정했다. 성별, 연령, 가구원 수, 혼인상태, 지역규모, 학력을 통제하고 오즈비와 95% 신뢰구간을 제시했다. 활동구성과 소비는 조사표의 당일형·숙박형 반복문항을 활동기록 단위로 변환해 가중 기술통계를 산출했다.")
    add_table(doc, ["지표", "산식·판정", "해석"], [
        ["체류·보급 필요도", "거리 백분위·시간 백분위·난이도 환산값의 평균", "이동부담이 클수록 증가"],
        ["숙박·보급·교통 접근성", "1km 내 존재=1, 1~3km=0.5, 3km 밖=0", "두 끝점 중 낮은 값 사용"],
        ["식음 접근성", "1km 3개 이상=1, 1~2개=0.75, 3km 내 3개 이상=0.5, 1~2개=0.25", "단순 존재보다 선택 가능성 반영"],
        ["기능별 공백", "필요도 × (1-기능별 접근성)", "이동부담 대비 부족한 기능"],
    ], widths=[3.3,7.7,5.6], font=7.5)
    add_body(doc, "구간 접근성은 시작점과 종료점 중 더 낮은 값을 적용했다. 식음·보급은 음식점과 보급시설 가운데 더 낮은 접근성을 사용한다. 따라서 두천1리처럼 편의점·식료품점은 가깝지만 일반 식사 제공 기능이 먼 경우, ‘기초 보급은 가능하지만 식사 기능은 부족한 구간’으로 해석된다.")
    add_body(doc, "지역 체류시장 여건지수는 방문규모, 숙박방문자 비율, 평균 체류시간, 외지인 관광소비를 0~1로 정규화해 평균했다. 이 값은 구간 공백점수에 합산하지 않고, 기존 사업자 연계와 경량 운영 중 어느 방식이 적절한지를 판단하는 보조지표로만 활용했다.")
    add_table(doc, ["검증 항목", "처리", "결과"], [
        ["GPX", "공식 시·종점 대조 및 9·55구간 방향 보정", "노선 연결 오류 방지"],
        ["시설 좌표", "좌표범위·영업상태·중복·행정경계 재검토", "상위구간 시설 접근성 재확인"],
        ["난이도", "대표값과 최고등급 기준 민감도 분석", "상위 우선순위의 안정성 확인"],
        ["교통", "정류장 위치와 실제 운행정보 분리", "51구간은 노선·배차 추가확인 후보로 제한"],
    ], widths=[3.0,7.0,6.6], font=7.5, header_fill="EAF2F8")
    add_source(doc, "공식 대피소·백패킹 지원은 상업숙박 공백 해석의 보조자료이며 핵심 공백점수에는 포함하지 않았다.")
    doc.add_page_break()

    # 5p
    add_title(doc, "4) 분석 내용 및 결과", "4-1. 숙박형 산림활동의 참여 특성과 서비스 요구")
    doc.add_picture(str(fig_dir / "figure1_activity_comparison.png"), width=Cm(16.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "그림 1. 당일형과 숙박형 산림활동의 주요 활동구성 비교", "자료: 통계청 MDIS 2025 산림휴양복지활동조사 마이크로데이터; 표본가중치 적용.")
    add_body(doc, "전체 응답자 11,949명 중 숙박형 활동 경험의 가중비율은 30.8%였다. 50대를 기준으로 숙박형 포함 오즈비는 20대 2.66, 30대 2.64, 40대 1.79로 높았고, 60대 0.76, 70세 이상 0.65로 낮았다. 이는 다른 인구특성을 통제한 관련성이며 연령이 숙박활동을 직접 유발한다는 뜻은 아니다.")
    add_body(doc, "숙박형은 당일형보다 캠핑 경험비율이 18.1%p, 수상활동이 11.7%p 높았고, 하이킹은 29.2%p 낮았다. 숙박형 산림활동은 당일 걷기를 단순 연장한 형태라기보다 캠핑과 체험이 결합된 활동구성에 가깝다.")
    day_spend = consumption.loc[consumption["유형"].eq("당일형"), "가중평균_만원"].iloc[0]
    overnight_spend = consumption.loc[consumption["유형"].eq("숙박형"), "가중평균_만원"].iloc[0]
    add_table(doc, ["핵심 결과", "당일형", "숙박형", "정책적 의미"], [
        ["1인 1회 평균 소비", f"{day_spend:.1f}만원", f"{overnight_spend:.1f}만원", "숙박형이 약 2.9배"],
        ["지역 선택", "교통 편의 비중 높음", "자연경관·먹거리·관광연계 비중 높음", "식사·체험 연계 필요"],
        ["주요 숙박", "-", "호텔·콘도 31.7%, 펜션 31.7%", "상업숙박 수요가 중심"],
        ["산촌마을·민박", "-", "4.7%", "신규 공급보다 예약·서비스 연계 우선"],
    ], widths=[3.6,3.5,5.0,4.3], font=7.4)
    doc.add_page_break()

    # 6p
    add_title(doc, "4) 분석 내용 및 결과", "4-2. 지역 체류시장 여건은 구간 순위가 아닌 보완방식을 가른다")
    doc.add_picture(str(fig_dir / "figure2_region_context.png"), width=Cm(16.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "그림 2. 분석대상 시·군의 체류시장 여건 보조지수", "자료: 한국관광 데이터랩, 2025.07~2026.06. 지수는 구간 공백점수와 분리해 사용.")
    rows = []
    for _, r in region.sort_values("지역체류시장여건지수", ascending=False).iterrows():
        rows.append([r["지역"], f"{r['방문자수_연인원_합계']/1e6:.1f}백만", f"{r['숙박방문자비율_월평균(%)']:.1f}%", f"{r['평균체류시간_월평균(분)']:.0f}분", f"{r['지역체류시장여건지수']:.1f}"])
    add_table(doc, ["지역", "12개월 방문", "숙박방문자 비율", "평균 체류시간", "보조지수"], rows, widths=[2.5,3.3,3.7,3.5,2.8], font=7.5)
    add_body(doc, "태안군은 방문규모와 숙박방문자 비율이 모두 높아 보조지수가 가장 높았다. 울진군은 방문규모는 작지만 숙박방문자 비율과 체류시간이 높아 체류형 운영의 가능성이 확인됐다. 반면 봉화군은 방문·소비 규모가 작아 대형 신규시설보다 대피소·소규모 보급·역과 정류장 연결을 묶는 경량 운영이 상대적으로 안전하다.")
    add_body(doc, "이 결과는 ‘관광객이 많으니 트레일 수요도 많다’는 뜻이 아니다. 같은 구간 공백이라도 태안처럼 기존 사업자가 많은 지역은 연계형 사업을, 봉화처럼 시장규모가 작은 지역은 예약형·소규모 실증을 우선할 수 있다는 보조 판단이다.")
    doc.add_page_break()

    # 7p
    add_title(doc, "4) 분석 내용 및 결과", "4-3. 전체 공개구간의 보완 우선순위")
    doc.add_picture(str(fig_dir / "figure3_overall_priority_map.png"), width=Cm(17.1))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "그림 3. 동서트레일 17개 공개구간의 보완 우선순위 지도", "자료: 숲나들e 공식 GPX·코스정보와 기능별 공백점수를 바탕으로 작성. 빨강=1~5위, 파랑=6~8위·비교구간.")
    add_body(doc, "상위 5개 가운데 54·53·51·47구간이 동부권에 위치했고, 서부권에서는 12구간의 상업숙박 공백이 확인됐다. 이는 지역 인지도보다 장거리 이동 후 도착하는 시·종점의 실제 기능 접근성이 우선순위를 좌우한다는 뜻이다.")
    add_body(doc, "55구간은 공식거리 20km로 이동부담이 가장 높은 수준이지만, 중섬마을과 망양정의 숙박·식음·교통 기반이 비교적 양호해 7위로 내려갔다. 길이가 긴 구간을 곧바로 시설투자 대상으로 지정하지 않는다는 점이 본 분석의 대표 비교사례다.")
    doc.add_page_break()

    # 8p
    add_title(doc, "4) 분석 내용 및 결과", "4-4. 상위 4개 구간의 취약지점과 실제 주변시설")
    doc.add_picture(str(fig_dir / "figure4_top4_detail_maps.png"), width=Cm(16.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "그림 4. 상위 4개 구간의 투영좌표 기반 1km·3km 접근권과 시설 위치", "자료: 공식 GPX, 숙박·상가·버스 공공데이터. 원은 EPSG:5174에서 산출한 실제 거리 버퍼이며 시설점은 검증된 최근접 시설.")
    add_body(doc, "54구간 두천1리 시작점은 음식점이 3km 안에 없고 최근접 음식점이 약 4.93km지만, 보급시설은 약 0.07km에 있다. 53구간 두천1리 종점도 음식점은 약 4.84km, 보급시설은 약 0.17km다. 두 구간은 ‘식음·보급 전체가 없는 곳’이 아니라 기초 보급은 가능하지만 실제 식사 기능이 부족한 구간이다.")
    add_body(doc, "51구간은 전곡리 종점에서 음식점 1.94km, 보급시설 2.38km, 정류장 약 2.00km로 나타났다. 보급지원형을 우선 적용하되 실제 노선과 배차는 추가 확인해야 한다. 12구간은 상업숙박이 약 4.15km로 멀지만 공식 대피소가 있어 신규 숙박시설보다 민박·대피소·식사·샤워를 묶는 전환형 운영이 적절하다.")
    doc.add_page_break()

    # 9p
    add_title(doc, "4) 분석 내용 및 결과", "4-5. 필요도와 기능별 공백을 분리한 운영유형")
    doc.add_picture(str(fig_dir / "figure5_need_gap_typology.png"), width=Cm(16.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "그림 5. 이동부담 기반 필요도, 기능별 공백과 권장 운영유형", "주: 필요도와 공백은 서로 다른 개념이므로 막대와 히트맵으로 분리해 표시. 55구간은 비교사례.")
    top = priority.head(6)
    rows=[]
    for _, r in top.iterrows():
        note = "노선·배차 추가확인" if int(r["구간번호"]) == 51 else ""
        rows.append([int(r["보완우선순위"]), f"{int(r['구간번호'])}구간", r["가장큰공백"], f"{r['최대공백점수']:.1f}", r["추천운영유형"], note])
    add_table(doc, ["순위", "구간", "가장 큰 공백", "점수", "운영유형", "검증사항"], rows, widths=[1.1,1.6,2.3,1.5,5.6,4.0], font=7.2)
    add_body(doc, "상위권의 중심은 숙박보다 식사·보급에 있다. 54·53·51·47구간은 보급지원형, 12구간은 상업숙박·마을서비스 전환형으로 분류됐다. 같은 보급지원형도 54구간은 출발 전 식사, 53구간은 기존 마을운영과 식사예약, 51구간은 종점 보급과 귀환수단 확인이라는 차이가 있다.")
    doc.add_page_break()

    # 10p
    add_title(doc, "5) 주요 성과 및 기대효과")
    add_table(doc, ["기존 접근", "본 분석의 변화"], [
        ["시설 수가 적은 지역을 막연히 취약하다고 판단", "이동부담에 비해 어느 기능이 비어 있는지를 구간 양 끝점에서 비교"],
        ["숙박·식당·교통을 한꺼번에 신규 설치", "보급지원·숙박전환·기존자원연계 등 구간별 최소 기능을 우선 적용"],
        ["지역 관광규모를 트레일 수요로 직접 해석", "개인·구간·시·군 자료를 분리하고 지역지표는 보조판단에만 사용"],
        ["정적 시설배치 계획", "개통 전 진단→시범운영→성과자료 반영→재진단의 환류체계"],
    ], widths=[7.7,8.8], font=7.8)
    add_body(doc, "정책적으로는 공개구간 확대 전에 상위구간을 먼저 점검하고, 거점마을 지원사업을 숙박·식사·보급·교통의 선택형 메뉴로 설계할 수 있다. 예약숲길에서는 사전예약 인원을 활용한 도시락·셔틀을, 자율트레킹에서는 영업시간과 정류장 정보를 표준화하는 등 운영방식을 달리할 수 있다.")
    add_body(doc, "숙박형 산림활동의 1인 평균 소비가 당일형보다 약 2.9배 높게 나타난 만큼 체류를 막는 기능 공백을 줄이는 것은 지역소비 연결 가능성을 높이는 방향이다. 다만 본 분석은 동서트레일 실제 이용자의 소비를 추적하지 않았으므로 매출 증가액이나 인과효과를 제시하지 않는다. 향후 예약률, 지역업체 이용건수, 식사·셔틀 이용률, 민원과 재방문 의향을 성과지표로 확인해야 한다.")
    add_callout(doc, "핵심 성과", "‘긴 구간=우선투자’가 아니라, 이동부담과 양 끝점의 서비스 접근성을 함께 진단해 구간마다 필요한 최소 기능을 다르게 배분하는 기준을 제시했다.", fill="E2F0D9")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    rr = p.add_run("참고자료")
    rr.bold = True
    rr.font.size = Pt(8.2)
    # Compact references, full URLs remain in the package CSV.
    for i, r in refs.iterrows():
        text = f"[{i+1}] {r['기관/저자']}, 「{r['자료명']}」, {r['일자/연도']}."
        pp = doc.add_paragraph()
        pp.paragraph_format.space_after = Pt(0)
        pp.paragraph_format.line_spacing = 1.0
        run = pp.add_run(text)
        run.font.size = Pt(6.5)
    add_source(doc, "참고자료의 URL과 활용내용은 제출 패키지 inputs/reference_list.csv에 수록했다.")

    docx = report_dir / "동서트레일_최종보고서_10p.docx"
    doc.save(docx)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(report_dir), str(docx)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    pdf = report_dir / "동서트레일_최종보고서_10p.pdf"
    return docx, pdf


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--csv-dir", type=Path, required=True)
    parser.add_argument("--fig-dir", type=Path, required=True)
    parser.add_argument("--inputs-dir", type=Path, required=True)
    parser.add_argument("--report-dir", type=Path, required=True)
    args = parser.parse_args()
    build(args.csv_dir, args.fig_dir, args.inputs_dir, args.report_dir)
